"""Parser for the BESS product database workbook (BESS_DATABASE.xlsx).

Every sheet in the workbook describes one container option (energy per
container + C-rate).  The sheets were prepared by hand and do not share an
identical layout, so the parser locates each block by its label instead of by
fixed cell addresses.  Anything the workbook does not provide is reported in
``Container.notes`` and ``Container.ac_source`` so the UI can flag it.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import BinaryIO, Union

from openpyxl import load_workbook

# AC-side chain in the order given in section 4 / 6 of the specification.
AC_KEYS = [
    "lv_cable",
    "pcs",
    "pcs_idt_cable",
    "idt",
    "dc_db",
    "mv_cable_switchgear",
    "main_transformer",
    "cable_11kv",
    "measurement",
]

AC_LABELS = {
    "lv_cable": "LV cable efficiency (DC cable)",
    "pcs": "PCS efficiency",
    "pcs_idt_cable": "PCS-to-IDT (LV) cable efficiency",
    "idt": "IDT efficiency",
    "dc_db": "DC DB efficiency",
    "mv_cable_switchgear": "MV cable & MV switchgear efficiency",
    "main_transformer": "Main power transformer efficiency",
    "cable_11kv": "11 kV cable efficiency (interconnecting point)",
    "measurement": "Measurement accuracy",
}


@dataclass
class Container:
    id: str
    sheet: str
    energy_mwh: float
    c_rate: float
    cycles_per_day: float
    aux_discharge_mw: float | None
    aux_charge_mw: float | None
    aux_standby_mw: float | None
    temp_labels: list[str]
    temp_upper: list[float]
    cal_rows: list[list[float]]  # cal_rows[m-1][band] = degradation in month m
    soh: list[float]  # index = year (0 = BOL)
    dc_rte: list[float]  # index = year (0 = BOL)
    ac: dict[str, float]
    ac_source: dict[str, str]
    ac_alt: dict[str, float | None]
    raw_ac: dict[str, float]
    notes: list[str] = field(default_factory=list)

    @property
    def label(self) -> str:
        return f"{self.energy_mwh:g} MWh @ {self.c_rate:g}C"

    @property
    def rated_power_mw(self) -> float:
        return self.energy_mwh * self.c_rate

    @property
    def aux_suspect(self) -> bool:
        """True when the aux loads exceed 20 % of rated power (values look like kW, not MW)."""
        vals = [v for v in (self.aux_discharge_mw, self.aux_charge_mw) if v is not None]
        return any(v > 0.2 * self.rated_power_mw for v in vals)

    @property
    def max_years(self) -> int:
        return min(len(self.soh), len(self.dc_rte)) - 1


# --------------------------------------------------------------------------
# low-level helpers
# --------------------------------------------------------------------------
def _num(value) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        m = re.search(r"-?\d+(?:\.\d+)?", value.replace(",", ""))
        return float(m.group()) if m else None
    return None


def _cells(ws) -> dict[tuple[int, int], object]:
    out = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and not (isinstance(c.value, str) and not c.value.strip()):
                out[(c.row, c.column)] = c.value
    return out


def _find(cells, predicate):
    return [
        (r, c)
        for (r, c), v in sorted(cells.items())
        if isinstance(v, str) and predicate(v.strip())
    ]


def _param(cells, pattern: str) -> float | None:
    """Value sitting to the right of the first label matching ``pattern``."""
    for (r, c), v in sorted(cells.items()):
        if isinstance(v, str) and re.search(pattern, v.strip(), re.I):
            for dc in (1, 2, 3):
                val = _num(cells.get((r, c + dc)))
                if val is not None:
                    return val
    return None


def _get(raw: dict[str, float], *patterns: str, exclude: str | None = None):
    for label, val in raw.items():
        low = label.lower()
        if all(re.search(p, low) for p in patterns) and not (exclude and re.search(exclude, low)):
            return label, val
    return None, None


# --------------------------------------------------------------------------
# AC chain normalisation
# --------------------------------------------------------------------------
def normalise_ac(raw: dict[str, float]):
    """Map the sheet-specific labels onto the nine AC-chain parameters.

    Mapping follows section 4 of the specification:
      * LV cable and PCS-to-IDT cable  <- "LV DC/AC cable efficiency"
      * PCS                            <- "PCS efficiency - discharging"
      * IDT and main power transformer <- "Transformer efficiency - discharging"
      * 11 kV cable                    <- "HV cable"
      * Measurement accuracy           <- "Measurement error"
    Parameters the sheet does not contain are set to 100 % and flagged.
    """
    lv_l, lv = _get(raw, r"^lv")
    pcs_l, pcs = _get(raw, r"^pcs", r"efficiency", exclude=r"idt")
    pi_l, pi = _get(raw, r"pcs", r"idt")
    idt_l, idt = _get(raw, r"^idt")
    tr_l, tr = _get(raw, r"^transformer")
    db_l, db = _get(raw, r"dc db")
    mv_l, mv = _get(raw, r"^mv cable")
    hvt_l, hvt = _get(raw, r"^hv transformer")
    hvc_l, hvc = _get(raw, r"^hv cable")
    me_l, me = _get(raw, r"measurement")

    missing_core = [n for n, v in [("LV cable", lv), ("PCS", pcs), ("MV cable", mv), ("Measurement", me)] if v is None]
    if missing_core:
        raise ValueError(f"AC efficiency table is missing: {', '.join(missing_core)}")

    ac: dict[str, float] = {}
    src: dict[str, str] = {}

    def put(key, label, value, how="database"):
        if value is None:
            ac[key] = 1.0
            src[key] = "not in database - 100 % assumed (editable)"
        else:
            ac[key] = float(value)
            src[key] = f"database: '{label}'" if how == "database" else f"derived from '{label}'"

    put("lv_cable", lv_l, lv)
    put("pcs", pcs_l, pcs)
    if pi is not None:
        put("pcs_idt_cable", pi_l, pi)
    else:
        put("pcs_idt_cable", lv_l, lv, "derived")
    if idt is not None:
        put("idt", idt_l, idt)
    else:
        put("idt", tr_l, tr, "derived")
    put("dc_db", db_l, db)
    put("mv_cable_switchgear", mv_l, mv)
    # Section 4: main power transformer <- "Transformer Efficiency - Discharging".
    # Sheets that do not carry that row (only an IDT efficiency) have no value.
    if tr is not None:
        put("main_transformer", tr_l, tr)
    else:
        put("main_transformer", None, None)
    if hvc is not None:
        put("cable_11kv", hvc_l, hvc)
    else:
        put("cable_11kv", None, None)
    put("measurement", me_l, me)

    alt = {"hv_transformer": hvt}
    return ac, src, alt


# --------------------------------------------------------------------------
# sheet parser
# --------------------------------------------------------------------------
def parse_sheet(ws) -> Container | None:
    cells = _cells(ws)
    energy = _param(cells, r"dc energy")
    if energy is None:
        return None
    notes: list[str] = []

    c_rate = _param(cells, r"^c[\s-]*rate")
    if c_rate is None:
        raise ValueError(f"Sheet '{ws.title}': C-rate not found")
    cycles = _param(cells, r"cycles") or 1.0
    aux_dis = _param(cells, r"aux.*dis\s*charg")
    aux_ch = _param(cells, r"aux.*(?<!dis)charge")
    aux_sb = _param(cells, r"aux.*standby")

    # --- calendar degradation block --------------------------------------
    pos = _find(cells, lambda s: "full month" in s.lower())
    if not pos:
        raise ValueError(f"Sheet '{ws.title}': calendar degradation table not found")
    r0, c0 = pos[0]
    bands, c = [], c0 + 1
    while isinstance(cells.get((r0, c)), str):
        bands.append(cells[(r0, c)].strip())
        c += 1
    uppers = []
    for b in bands:
        nums = [float(x) for x in re.findall(r"\d+(?:\.\d+)?", b)]
        uppers.append(max(nums))
    rows, r = [], r0 + 1
    while isinstance(cells.get((r, c0)), (int, float)):
        rows.append([float(cells[(r, c0 + 1 + i)]) for i in range(len(bands))])
        r += 1

    # --- SOH / DC RTE block -------------------------------------------------
    pos = _find(cells, lambda s: s.lower() == "year")
    if not pos:
        raise ValueError(f"Sheet '{ws.title}': SOH / DC RTE table not found")
    r0, yc = pos[0]
    soh_c = next((c for c in range(yc + 1, yc + 6) if "soh" in str(cells.get((r0, c), "")).lower()), None)
    rte_c = next((c for c in range(yc + 1, yc + 6) if "rte" in str(cells.get((r0, c), "")).lower()), None)
    if soh_c is None or rte_c is None:
        raise ValueError(f"Sheet '{ws.title}': SOH / DC RTE columns not found")
    soh, rte, r = [], [], r0 + 1
    while isinstance(cells.get((r, yc)), (int, float)):
        soh.append(float(cells[(r, soh_c)]))
        rte.append(float(cells[(r, rte_c)]))
        r += 1

    # --- AC efficiency block ---------------------------------------------------
    raw: dict[str, float] = {}
    for (r0, c0) in _find(cells, lambda s: s.lower() == "parameter"):
        hdr = cells.get((r0, c0 + 1))
        if isinstance(hdr, str) and "efficiency" in hdr.lower():
            r = r0 + 1
            while isinstance(cells.get((r, c0)), str) and _num(cells.get((r, c0 + 1))) is not None:
                raw[cells[(r, c0)].strip()] = float(_num(cells[(r, c0 + 1)]))
                r += 1
            break
    if not raw:
        raise ValueError(f"Sheet '{ws.title}': AC efficiency table not found")
    ac, src, alt = normalise_ac(raw)

    # --- data-quality notes -------------------------------------------------------
    power = energy * c_rate
    if aux_dis is None or aux_ch is None:
        notes.append("Auxiliary load per container is not in the database - enter it in 'Override database values'.")
    else:
        for lab, val in (("discharge", aux_dis), ("charge", aux_ch)):
            if val > 0.2 * power:
                notes.append(
                    f"Aux load ({lab}) is {val:g} MW = {val / power:.0%} of the container's rated power "
                    f"({power:.2f} MW). The value looks like kW rather than MW - please verify the units."
                )
    for k, v in src.items():
        if v.startswith("not in database"):
            notes.append(f"{AC_LABELS[k]}: not in database, 100 % assumed.")
    if len(soh) != len(rte):
        notes.append("SOH and DC RTE tables have different lengths.")

    return Container(
        id=f"{energy:g}MWh_{c_rate:g}C",
        sheet=ws.title,
        energy_mwh=float(energy),
        c_rate=float(c_rate),
        cycles_per_day=float(cycles),
        aux_discharge_mw=aux_dis,
        aux_charge_mw=aux_ch,
        aux_standby_mw=aux_sb,
        temp_labels=bands,
        temp_upper=uppers,
        cal_rows=rows,
        soh=soh,
        dc_rte=rte,
        ac=ac,
        ac_source=src,
        ac_alt=alt,
        raw_ac=raw,
        notes=notes,
    )


def load_database(source: Union[str, Path, BinaryIO]) -> list[Container]:
    """Parse every container sheet of the workbook."""
    wb = load_workbook(source, data_only=True)
    containers: list[Container] = []
    for ws in wb.worksheets:
        cont = parse_sheet(ws)
        if cont is not None:
            containers.append(cont)
    if not containers:
        raise ValueError("No container sheets found in the workbook")
    ids = [c.id for c in containers]
    if len(set(ids)) != len(ids):
        for i, c in enumerate(containers):
            c.id = f"{c.id}_{i + 1}"
    return containers


def select_by_c_rate(containers: list[Container], c_rate: float) -> list[Container]:
    """Containers whose C-rate is nearest to ``c_rate`` (all containers on a tie)."""
    best = min(abs(c.c_rate - c_rate) for c in containers)
    return [c for c in containers if abs(abs(c.c_rate - c_rate) - best) < 1e-9]
