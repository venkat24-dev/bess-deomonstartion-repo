"""Excel export of the sizing results (one block of sheets per container option)."""
from __future__ import annotations

import io

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .engine import SizingOutcome, status_labels

GREEN = PatternFill("solid", start_color="C6EFCE", end_color="C6EFCE")
RED = PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")
GOOD = {"TRUE", "MATCHED"}
BAD = {"FALSE", "NOT MATCHED"}


def verification_frame(res, style: str) -> pd.DataFrame:
    v = res.verify.copy()
    out = v[
        [
            "Required energy (MWh)",
            "Calculated energy (MWh)",
            "Required RTE (%)",
            "Calculated RTE incl. aux (%)",
            "Required power (MW)",
            "Calculated power (MW)",
        ]
    ].copy()
    out["Energy"] = status_labels(v["Energy OK"], style)
    out["RTE"] = status_labels(v["RTE OK"], style)
    out["Power"] = status_labels(v["Power OK"], style)
    out["Overall"] = status_labels(v["All OK"], style)
    return out


def build_excel(outcomes: list[SizingOutcome], style: str = "MATCHED") -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        rows = []
        for o in outcomes:
            s = o.result.summary
            rows.append(
                {
                    "Container": s["container"],
                    "Initial containers": s["n_init"],
                    "Installed DC capacity (MWh)": s["installed_mwh"],
                    "Margin applied (%)": o.margin_applied_pct,
                    "Effective margin (%)": s["effective_margin_pct"],
                    "BOL energy at POI (MWh)": s["bol_energy"],
                    "EOL energy at POI (MWh)": s["eol_energy"],
                    "EOL RTE incl. aux (%)": s["rte_eol_incl"],
                    "Cycles OK": s["cycles_ok"],
                    "Accepted": s["accepted"],
                    "Note": o.note,
                }
            )
        pd.DataFrame(rows).to_excel(xw, sheet_name="Summary", index=False)
        for i, o in enumerate(outcomes, 1):
            r = o.result
            tag = f"Opt{i}"
            verification_frame(r, style).to_excel(xw, sheet_name=f"{tag} Verification")
            r.dis.to_excel(xw, sheet_name=f"{tag} Discharging")
            r.chg.to_excel(xw, sheet_name=f"{tag} Charging")
            r.perf.to_excel(xw, sheet_name=f"{tag} RTE-Power")
        for ws in xw.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for idx, col in enumerate(ws.columns, 1):
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[get_column_letter(idx)].width = min(48, max(10, width + 2))
            if ws.title.endswith("Verification"):
                for row in ws.iter_rows(min_row=2):
                    for c in row:
                        if c.value in GOOD:
                            c.fill = GREEN
                        elif c.value in BAD:
                            c.fill = RED
    return buf.getvalue()
