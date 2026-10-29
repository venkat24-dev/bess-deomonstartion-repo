"""Validated, explicit import of the Excel container catalog."""
from __future__ import annotations

import hashlib
import io
import math
import re
from collections import Counter
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Any
from uuid import UUID
from zipfile import BadZipFile

from openpyxl.utils.exceptions import InvalidFileException
from openpyxl import load_workbook
from sqlalchemy import select, text, update
from sqlalchemy.orm import Session, selectinload

from bess.database import AC_KEYS, Container, _num, load_database
from bess.models import (
    CatalogRevision,
    ContainerACParameter,
    ContainerCalendarDegradation,
    ContainerModel,
    ContainerModelNote,
    ContainerPerformanceYear,
    ContainerRawACParameter,
    ContainerTemperatureBand,
)

MAX_CATALOG_UPLOAD_SIZE_MB = 200
MAX_CATALOG_UPLOAD_SIZE_BYTES = MAX_CATALOG_UPLOAD_SIZE_MB * 1024 * 1024


@dataclass(frozen=True)
class ContainerValidation:
    source_key: str
    temperature_band_count: int
    calendar_month_count: int
    performance_year_count: int
    raw_ac_parameter_count: int
    normalized_ac_parameter_count: int


@dataclass(frozen=True)
class CatalogValidationReport:
    source_filename: str
    source_sha256: str
    is_valid: bool
    containers: tuple[ContainerValidation, ...]
    errors: tuple[str, ...]
    warnings: tuple[str, ...]


@dataclass(frozen=True)
class CatalogImportResult:
    revision_id: UUID
    source_sha256: str
    container_count: int
    created: bool


class CatalogValidationError(ValueError):
    def __init__(self, report: CatalogValidationReport):
        self.report = report
        detail = "; ".join(report.errors) or "Catalog validation failed"
        super().__init__(detail)


class CatalogUploadError(ValueError):
    """The uploaded file is not an allowed catalog workbook."""


class CatalogIntegrityError(RuntimeError):
    """Stored rows do not match the immutable workbook revision."""


class CatalogRevisionNotFound(LookupError):
    """The requested catalog revision does not exist."""


@dataclass(frozen=True)
class _CatalogPlan:
    report: CatalogValidationReport
    parsed: tuple[Container, ...]


def _finite(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def _source_key(container: Container) -> str:
    return f"{container.energy_mwh:g}MWh_{container.c_rate:g}C"


def _read_plan(source: str | Path | bytes, source_filename: str | None = None) -> _CatalogPlan:
    if isinstance(source, bytes):
        payload = source
        filename = source_filename or "uploaded.xlsx"
    else:
        path = Path(source)
        filename = path.name
        try:
            payload = path.read_bytes()
        except OSError as exc:
            report = CatalogValidationReport(
                filename, "", False, (), (f"Cannot read workbook: {exc}",), ()
            )
            return _CatalogPlan(report, ())

    digest = hashlib.sha256(payload).hexdigest()
    errors: list[str] = []
    warnings: list[str] = []
    try:
        parsed = tuple(load_database(io.BytesIO(payload)))
        workbook = load_workbook(io.BytesIO(payload), data_only=True, read_only=True)
    except (BadZipFile, InvalidFileException, IndexError, KeyError, OSError, TypeError, ValueError) as exc:
        report = CatalogValidationReport(
            filename,
            digest,
            False,
            (),
            (f"Workbook could not be parsed: {exc}",),
            (),
        )
        return _CatalogPlan(report, ())

    keys = [_source_key(container) for container in parsed]
    duplicates = sorted(key for key, count in Counter(keys).items() if count > 1)
    if duplicates:
        errors.append(f"Duplicate container source keys: {', '.join(duplicates)}")
    parsed_sheets = {container.sheet for container in parsed}
    skipped_sheets = sorted(set(workbook.sheetnames) - parsed_sheets)
    if skipped_sheets:
        errors.append(
            "Workbook sheets were skipped because the parser found no container record: "
            + ", ".join(skipped_sheets)
        )

    summaries: list[ContainerValidation] = []
    for container in parsed:
        key = _source_key(container)
        prefix = f"{key} (sheet '{container.sheet}')"
        raw_cycles = _source_parameter(workbook[container.sheet], r"cycles")
        if raw_cycles is None or not _finite(raw_cycles) or raw_cycles <= 0:
            errors.append(f"{prefix}: source workbook must contain a finite positive cycles value")
        if container.id != key:
            errors.append(f"{prefix}: parser identifier {container.id!r} does not match canonical source key")
        for field_name, value in (
            ("energy_mwh", container.energy_mwh),
            ("c_rate", container.c_rate),
            ("cycles_per_day", container.cycles_per_day),
        ):
            if not _finite(value) or value <= 0:
                errors.append(f"{prefix}: {field_name} must be a finite positive number")

        band_count = len(container.temp_labels)
        if not band_count or len(container.temp_upper) != band_count:
            errors.append(f"{prefix}: temperature labels and upper bounds must be present and aligned")
        elif (
            any(not label.strip() for label in container.temp_labels)
            or len(set(container.temp_labels)) != band_count
            or any(not _finite(upper) for upper in container.temp_upper)
            or any(a >= b for a, b in zip(container.temp_upper, container.temp_upper[1:]))
        ):
            errors.append(f"{prefix}: temperature bands must have unique labels and increasing finite bounds")

        if not container.cal_rows:
            errors.append(f"{prefix}: calendar degradation rows are missing")
        for month, row in enumerate(container.cal_rows, start=1):
            if len(row) != band_count:
                errors.append(f"{prefix}: calendar degradation month {month} does not match temperature bands")
            if any(not _finite(value) or not 0 <= value < 1 for value in row):
                errors.append(f"{prefix}: calendar degradation month {month} contains an invalid value")

        if not container.soh or not container.dc_rte:
            errors.append(f"{prefix}: SOH and DC RTE performance years are required")
        if len(container.soh) != len(container.dc_rte):
            errors.append(f"{prefix}: SOH and DC RTE year counts differ")
        for name, values in (("SOH", container.soh), ("DC RTE", container.dc_rte)):
            if any(not _finite(value) or not 0 < value <= 1 for value in values):
                errors.append(f"{prefix}: {name} values must be finite and between 0 and 1")
        if container.soh and container.soh[0] != 1.0:
            errors.append(f"{prefix}: year-zero SOH must be 1.0")

        missing_ac_keys = set(AC_KEYS) - set(container.ac)
        extra_ac_keys = set(container.ac) - set(AC_KEYS)
        if missing_ac_keys or extra_ac_keys:
            errors.append(
                f"{prefix}: normalized AC keys mismatch; missing={sorted(missing_ac_keys)}, "
                f"unexpected={sorted(extra_ac_keys)}"
            )
        for ac_key in AC_KEYS:
            if ac_key in container.ac and (
                not _finite(container.ac[ac_key]) or not 0 < container.ac[ac_key] <= 1
            ):
                errors.append(f"{prefix}: normalized AC efficiency '{ac_key}' must be finite and in (0, 1]")
            provenance = container.ac_source.get(ac_key)
            if provenance is None or not (
                provenance.startswith("database:")
                or provenance.startswith("derived from ")
                or provenance.startswith("not in database")
            ):
                errors.append(f"{prefix}: normalized AC provenance '{ac_key}' is missing or invalid")

        hv_transformer = container.ac_alt.get("hv_transformer")
        if hv_transformer is not None and (
            not _finite(hv_transformer) or not 0 < hv_transformer <= 1
        ):
            errors.append(f"{prefix}: alternate HV transformer efficiency must be finite and in (0, 1]")
        for label, value in container.raw_ac.items():
            if not label.strip() or not _finite(value) or not 0 < value <= 1:
                errors.append(f"{prefix}: raw AC source '{label}' has an invalid efficiency")

        for aux_name, value in (
            ("discharge", container.aux_discharge_mw),
            ("charge", container.aux_charge_mw),
            ("standby", container.aux_standby_mw),
        ):
            if value is None:
                warnings.append(f"{prefix}: {aux_name} auxiliary load is missing; preserved as NULL.")
            elif not _finite(value) or value < 0:
                errors.append(f"{prefix}: {aux_name} auxiliary load must be a finite nonnegative number")

        warnings.extend(f"{prefix}: {note}" for note in container.notes)
        summaries.append(
            ContainerValidation(
                key,
                band_count,
                len(container.cal_rows),
                len(container.soh),
                len(container.raw_ac),
                len(container.ac),
            )
        )

    report = CatalogValidationReport(
        filename,
        digest,
        not errors,
        tuple(summaries),
        tuple(errors),
        tuple(warnings),
    )
    return _CatalogPlan(report, parsed)


def _source_parameter(worksheet: Any, pattern: str) -> float | None:
    """Verify fields whose parser has a fallback that hides missing source data."""
    for row in worksheet.iter_rows():
        for cell in row:
            if not isinstance(cell.value, str) or not re.search(pattern, cell.value, re.I):
                continue
            for offset in (1, 2, 3):
                value = _num(worksheet.cell(cell.row, cell.column + offset).value)
                if value is not None:
                    return value
    return None


def dry_run_catalog(source: str | Path) -> CatalogValidationReport:
    """Parse and validate a workbook without opening or writing to a database."""
    return _read_plan(source).report


def dry_run_catalog_bytes(payload: bytes, source_filename: str) -> CatalogValidationReport:
    """Parse and validate uploaded workbook bytes without writing to a database."""
    return _read_plan(payload, source_filename).report


def parse_catalog_upload(payload: bytes, source_filename: str) -> list[Container]:
    """Validate an uploaded XLSX workbook before returning its usable catalog."""
    if Path(source_filename).suffix.casefold() != ".xlsx":
        raise CatalogUploadError(
            "Unsupported product database file type. Upload an Excel .xlsx workbook; CSV files are not supported."
        )
    if len(payload) > MAX_CATALOG_UPLOAD_SIZE_BYTES:
        raise CatalogUploadError(
            f"Product database files must be no larger than {MAX_CATALOG_UPLOAD_SIZE_MB} MB."
        )

    plan = _read_plan(payload, source_filename)
    if not plan.report.is_valid:
        raise CatalogValidationError(plan.report)
    return list(plan.parsed)


def _decimal(value: float) -> Decimal:
    return Decimal(str(value))


def _provenance(source_text: str) -> tuple[str, str | None]:
    if source_text.startswith("database: '") and source_text.endswith("'"):
        return "database", source_text[len("database: '") : -1]
    if source_text.startswith("derived from '") and source_text.endswith("'"):
        return "derived", source_text[len("derived from '") : -1]
    if source_text.startswith("not in database"):
        return "assumed", None
    raise CatalogIntegrityError(f"Unsupported AC provenance value: {source_text!r}")


def _container_to_engine(model: ContainerModel) -> Container:
    years = sorted(model.performance_years, key=lambda row: row.year_index)
    bands = sorted(model.temperature_bands, key=lambda row: row.band_index)
    degradation_rows: dict[int, dict[int, float]] = {}
    for row in model.calendar_degradation:
        degradation_rows.setdefault(row.month_index, {})[row.band_index] = float(row.degradation)
    ac = {
        row.parameter_key: float(row.efficiency)
        for row in model.ac_parameters
        if row.parameter_key != "hv_transformer" and row.efficiency is not None
    }
    ac_rows = {row.parameter_key: row for row in model.ac_parameters}
    hv = ac_rows.get("hv_transformer")
    return Container(
        id=model.source_key,
        sheet=model.source_sheet,
        energy_mwh=float(model.energy_mwh),
        c_rate=float(model.c_rate),
        cycles_per_day=float(model.cycles_per_day),
        aux_discharge_mw=(
            float(model.aux_discharge_mw) if model.aux_discharge_mw is not None else None
        ),
        aux_charge_mw=float(model.aux_charge_mw) if model.aux_charge_mw is not None else None,
        aux_standby_mw=float(model.aux_standby_mw) if model.aux_standby_mw is not None else None,
        temp_labels=[row.label for row in bands],
        temp_upper=[float(row.upper_c) for row in bands],
        cal_rows=[
            [degradation_rows[month][band.band_index] for band in bands]
            for month in sorted(degradation_rows)
        ],
        soh=[float(row.soh) for row in years],
        dc_rte=[float(row.dc_rte) for row in years],
        ac=ac,
        ac_source={
            row.parameter_key: row.source_text or ""
            for row in model.ac_parameters
            if row.parameter_key != "hv_transformer"
        },
        ac_alt={"hv_transformer": float(hv.efficiency) if hv and hv.efficiency is not None else None},
        raw_ac={row.source_label: float(row.value) for row in model.raw_ac_parameters},
        notes=[row.note for row in sorted(model.data_notes, key=lambda note: note.note_index)],
    )


def _load_revision(session: Session, revision_id: UUID) -> CatalogRevision | None:
    statement = (
        select(CatalogRevision)
        .where(CatalogRevision.id == revision_id)
        .options(
            selectinload(CatalogRevision.container_models).selectinload(
                ContainerModel.performance_years
            ),
            selectinload(CatalogRevision.container_models).selectinload(
                ContainerModel.temperature_bands
            ),
            selectinload(CatalogRevision.container_models).selectinload(
                ContainerModel.calendar_degradation
            ),
            selectinload(CatalogRevision.container_models).selectinload(
                ContainerModel.ac_parameters
            ),
            selectinload(CatalogRevision.container_models).selectinload(
                ContainerModel.raw_ac_parameters
            ),
            selectinload(CatalogRevision.container_models).selectinload(
                ContainerModel.data_notes
            ),
        )
    )
    return session.scalar(statement)


def load_catalog_revision(revision_id: UUID, session: Session) -> list[Container]:
    """Reconstruct engine inputs from one persisted catalog revision."""
    revision = _load_revision(session, revision_id)
    if revision is None:
        raise CatalogRevisionNotFound(f"Catalog revision {revision_id} does not exist.")
    return [
        _container_to_engine(model)
        for model in sorted(revision.container_models, key=lambda item: item.source_key)
    ]


def _stored_revision_matches(
    revision: CatalogRevision, parsed: tuple[Container, ...]
) -> bool:
    try:
        stored = {
            model.source_key: _container_to_engine(model)
            for model in revision.container_models
        }
    except (IndexError, KeyError, TypeError, ValueError):
        return False
    expected = {_source_key(container): container for container in parsed}
    return stored == expected


def _persist_container(session: Session, revision_id: UUID, container: Container) -> None:
    source_key = _source_key(container)
    model = ContainerModel(
        catalog_revision_id=revision_id,
        source_key=source_key,
        source_sheet=container.sheet,
        energy_mwh=_decimal(container.energy_mwh),
        c_rate=_decimal(container.c_rate),
        cycles_per_day=_decimal(container.cycles_per_day),
        aux_discharge_mw=(
            _decimal(container.aux_discharge_mw)
            if container.aux_discharge_mw is not None
            else None
        ),
        aux_charge_mw=(
            _decimal(container.aux_charge_mw) if container.aux_charge_mw is not None else None
        ),
        aux_standby_mw=(
            _decimal(container.aux_standby_mw) if container.aux_standby_mw is not None else None
        ),
    )
    session.add(model)
    session.flush()

    for year, (soh, dc_rte) in enumerate(zip(container.soh, container.dc_rte)):
        session.add(
            ContainerPerformanceYear(
                container_model_id=model.id,
                year_index=year,
                soh=_decimal(soh),
                dc_rte=_decimal(dc_rte),
            )
        )
    for band_index, (label, upper) in enumerate(zip(container.temp_labels, container.temp_upper)):
        session.add(
            ContainerTemperatureBand(
                container_model_id=model.id,
                band_index=band_index,
                label=label,
                upper_c=_decimal(upper),
            )
        )
    session.flush()
    for month_index, values in enumerate(container.cal_rows, start=1):
        for band_index, value in enumerate(values):
            session.add(
                ContainerCalendarDegradation(
                    container_model_id=model.id,
                    month_index=month_index,
                    band_index=band_index,
                    degradation=_decimal(value),
                )
            )

    for parameter_key in AC_KEYS:
        source_text = container.ac_source[parameter_key]
        source_kind, source_label = _provenance(source_text)
        session.add(
            ContainerACParameter(
                container_model_id=model.id,
                parameter_key=parameter_key,
                efficiency=_decimal(container.ac[parameter_key]),
                source_kind=source_kind,
                source_label=source_label,
                source_text=source_text,
            )
        )
    hv_transformer = container.ac_alt.get("hv_transformer")
    hv_label = next(
        (label for label in container.raw_ac if re.search(r"^hv transformer", label, re.I)),
        None,
    )
    session.add(
        ContainerACParameter(
            container_model_id=model.id,
            parameter_key="hv_transformer",
            efficiency=_decimal(hv_transformer) if hv_transformer is not None else None,
            source_kind="database" if hv_transformer is not None else "missing",
            source_label=hv_label,
        )
    )
    for source_label, value in container.raw_ac.items():
        session.add(
            ContainerRawACParameter(
                container_model_id=model.id,
                source_label=source_label,
                value=_decimal(value),
            )
        )
    for note_index, note in enumerate(container.notes):
        session.add(
            ContainerModelNote(
                container_model_id=model.id,
                note_index=note_index,
                note=note,
            )
        )


def import_catalog(source: str | Path, session: Session) -> CatalogImportResult:
    """Persist a validated revision atomically; this never activates it."""
    plan = _read_plan(source)
    if not plan.report.is_valid:
        raise CatalogValidationError(plan.report)

    with session.begin():
        revision = session.scalar(
            select(CatalogRevision).where(
                CatalogRevision.source_sha256 == plan.report.source_sha256
            )
        )
        if revision is not None:
            loaded = _load_revision(session, revision.id)
            if loaded is None or not _stored_revision_matches(loaded, plan.parsed):
                raise CatalogIntegrityError(
                    "An existing revision with this workbook hash does not match the parsed source."
                )
            return CatalogImportResult(
                revision.id,
                revision.source_sha256,
                len(plan.parsed),
                False,
            )

        revision = CatalogRevision(
            source_filename=plan.report.source_filename,
            source_sha256=plan.report.source_sha256,
            is_active=False,
        )
        session.add(revision)
        session.flush()
        for container in plan.parsed:
            _persist_container(session, revision.id, container)
        session.flush()
        return CatalogImportResult(
            revision.id,
            revision.source_sha256,
            len(plan.parsed),
            True,
        )


def activate_catalog_revision(revision_id: UUID, session: Session) -> None:
    """Explicitly activate one revision, serializing concurrent activations."""
    with session.begin():
        if session.connection().dialect.name == "postgresql":
            session.execute(
                text("SELECT pg_advisory_xact_lock(:namespace, :lock_id)"),
                {"namespace": 1111573586, "lock_id": 1},
            )
        revision = session.scalar(
            select(CatalogRevision)
            .where(CatalogRevision.id == revision_id)
            .with_for_update()
        )
        if revision is None:
            raise CatalogRevisionNotFound(f"Catalog revision {revision_id} does not exist.")
        session.execute(
            update(CatalogRevision)
            .where(CatalogRevision.is_active.is_(True))
            .values(is_active=False)
        )
        revision.is_active = True
