from __future__ import annotations

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from bess import catalog_import
from bess.catalog_import import (
    MAX_CATALOG_UPLOAD_SIZE_BYTES,
    CatalogUploadError,
    CatalogValidationError,
    parse_catalog_upload,
)
from bess.database import load_database
from bess.engine import Inputs, compute_case, effective_params, status_labels


PROJECT_ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = PROJECT_ROOT / "data" / "BESS_DATABASE.xlsx"


def _workbook_bytes(tmp_path, edit):
    path = tmp_path / "catalog.xlsx"
    workbook = load_workbook(WORKBOOK)
    edit(workbook)
    workbook.save(path)
    return path.read_bytes()


def _cell_right_of(sheet, label):
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and label.casefold() in cell.value.casefold():
                for offset in (1, 2, 3):
                    adjacent = sheet.cell(cell.row, cell.column + offset)
                    if adjacent.value is not None:
                        return adjacent
    raise AssertionError(f"Could not find workbook value for {label!r}")


def _replace_label(workbook, label, replacement):
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.casefold() == label.casefold():
                    cell.value = replacement
                    return
    raise AssertionError(f"Could not find workbook label {label!r}")


def test_upload_parser_accepts_only_validated_xlsx_and_returns_all_options():
    containers = parse_catalog_upload(WORKBOOK.read_bytes(), WORKBOOK.name)

    assert len(containers) == 6
    assert {container.id for container in containers} == {
        container.id for container in load_database(WORKBOOK)
    }


@pytest.mark.parametrize(
    ("filename", "payload", "message"),
    [
        ("products.csv", b"energy,c_rate\n1,0.5", "CSV files are not supported"),
        ("products.xls", b"not an xlsx file", "Upload an Excel .xlsx workbook"),
    ],
)
def test_upload_parser_rejects_unsupported_file_types(filename, payload, message):
    with pytest.raises(CatalogUploadError, match=message):
        parse_catalog_upload(payload, filename)


def test_upload_parser_rejects_oversized_files(monkeypatch):
    monkeypatch.setattr(catalog_import, "MAX_CATALOG_UPLOAD_SIZE_BYTES", 4)

    with pytest.raises(CatalogUploadError, match="no larger than 200 MB"):
        parse_catalog_upload(b"12345", "large.xlsx")

    assert MAX_CATALOG_UPLOAD_SIZE_BYTES == 200 * 1024 * 1024


@pytest.mark.parametrize(
    ("edit", "message"),
    [
        (
            lambda workbook: _cell_right_of(workbook["7.044 MWh 0.25C"], "DC Energy per Container").__setattr__(
                "value", "bad capacity"
            ),
            "Workbook sheets were skipped|could not be parsed",
        ),
        (
            lambda workbook: _cell_right_of(workbook["7.044 MWh 0.25C"], "C Rate").__setattr__(
                "value", "bad C-rate"
            ),
            "C-rate not found",
        ),
        (
            lambda workbook: _replace_label(
                workbook, "Full Month Between FAT and SAT", "removed"
            ),
            "calendar degradation table not found",
        ),
    ],
)
def test_upload_parser_rejects_missing_or_incorrect_required_values(tmp_path, edit, message):
    payload = _workbook_bytes(tmp_path, edit)

    with pytest.raises(CatalogValidationError, match=message):
        parse_catalog_upload(payload, "invalid.xlsx")


def test_upload_parser_rejects_duplicate_products(tmp_path):
    def make_duplicate(workbook):
        first = workbook["7.044 MWh 0.25C"]
        _cell_right_of(first, "DC Energy per Container").value = "5.015 MWh"
        _cell_right_of(first, "C Rate").value = "0.25 C"

    payload = _workbook_bytes(tmp_path, make_duplicate)

    with pytest.raises(CatalogValidationError, match="Duplicate container source keys"):
        parse_catalog_upload(payload, "duplicate.xlsx")


def test_upload_parser_rejects_negative_capacity_and_auxiliary_values(tmp_path):
    def make_negative(workbook):
        sheet = workbook["7.044 MWh 0.25C"]
        _cell_right_of(sheet, "DC Energy per Container").value = "-7.044 MWh"

    with pytest.raises(CatalogValidationError, match="energy_mwh must be a finite positive number"):
        parse_catalog_upload(_workbook_bytes(tmp_path, make_negative), "negative-capacity.xlsx")

    def make_negative_aux(workbook):
        sheet = next(sheet for sheet in workbook.worksheets if sheet.title.startswith("6.626 MWh 0.5"))
        _cell_right_of(sheet, "Aux Load - Discharge").value = -2.25

    with pytest.raises(CatalogValidationError, match="auxiliary load must be a finite nonnegative number"):
        parse_catalog_upload(_workbook_bytes(tmp_path, make_negative_aux), "negative-aux.xlsx")


@pytest.mark.parametrize("payload", [b"", b"not an Excel workbook"])
def test_upload_parser_rejects_empty_and_corrupted_workbooks(payload):
    with pytest.raises(CatalogValidationError):
        parse_catalog_upload(payload, "corrupt.xlsx")


def test_match_labels_change_without_changing_the_boolean_result():
    matched = pd.Series([True, False])

    assert status_labels(matched).tolist() == ["MATCHED", "NOT MATCHED"]
    assert status_labels(matched, "TRUE/FALSE").tolist() == ["TRUE", "FALSE"]
    assert matched.tolist() == [True, False]


def test_auxiliary_kilowatt_setting_changes_only_flagged_product_outputs():
    database = load_database(WORKBOOK)
    flagged = next(container for container in database if container.id == "6.626MWh_0.5C")
    unflagged = next(container for container in database if container.id == "5.015MWh_0.25C")

    raw = effective_params(flagged, Inputs(aux_kw_fix=False))
    corrected = effective_params(flagged, Inputs(aux_kw_fix=True))
    assert corrected["aux_dis"] == pytest.approx(raw["aux_dis"] / 1000)
    assert corrected["aux_ch"] == pytest.approx(raw["aux_ch"] / 1000)
    assert corrected["aux_sb"] == pytest.approx(raw["aux_sb"] / 1000)
    assert effective_params(unflagged, Inputs(aux_kw_fix=True)) == effective_params(
        unflagged, Inputs(aux_kw_fix=False)
    )

    without_fix = compute_case(flagged, 1, Inputs(aux_kw_fix=False))
    with_fix = compute_case(flagged, 1, Inputs(aux_kw_fix=True))
    assert without_fix.dis.loc[0, "Aux energy (MWh)"] == pytest.approx(9.0)
    assert with_fix.dis.loc[0, "Aux energy (MWh)"] == pytest.approx(0.009)
    assert with_fix.perf.loc[0, "Net power at POI (MW)"] > without_fix.perf.loc[
        0, "Net power at POI (MW)"
    ]


def test_settings_upload_updates_sizing_and_recovers_after_invalid_upload(tmp_path):
    from streamlit.testing.v1 import AppTest

    def summary_frame(app):
        return next(
            widget.value
            for widget in app.dataframe
            if isinstance(widget.value, pd.DataFrame)
            and "Container" in widget.value.columns
            and "Result" in widget.value.columns
        )

    def auxiliary_override_frame(app):
        return next(
            widget.value
            for widget in app.dataframe
            if isinstance(widget.value, pd.DataFrame)
            and "Parameter" in widget.value.columns
            and "Value" in widget.value.columns
            and widget.value["Parameter"].astype(str).str.contains(
                "Aux load - discharge"
            ).any()
        )

    def performance_frame(app):
        return next(
            widget.value
            for widget in app.dataframe
            if isinstance(widget.value, pd.DataFrame)
            and "Net power at POI (MW)" in widget.value.columns
        )

    custom_workbook = _workbook_bytes(
        tmp_path,
        lambda workbook: _cell_right_of(
            workbook["7.044 MWh 0.25C"], "DC Energy per Container"
        ).__setattr__("value", "8.0 MWh"),
    )
    at = AppTest.from_file(str(PROJECT_ROOT / "app.py"), default_timeout=120).run()
    assert not at.exception
    assert any("Bundled database: 6 container options" in item.value for item in at.caption)
    assert any("7.044 MWh @ 0.25C" in item.value for item in at.info)
    assert any("34.0 %" in item.value for item in at.caption)

    summary_before = summary_frame(at)
    labels_before = summary_before["Result"].tolist()
    numeric_before = summary_before.drop(columns=["Result"])
    next(widget for widget in at.radio if widget.label == "Match result labels").set_value(
        "TRUE / FALSE"
    ).run()
    summary_after = summary_frame(at)
    assert summary_after["Result"].tolist() == [
        "TRUE" if value == "MATCHED" else "FALSE" for value in labels_before
    ]
    pd.testing.assert_frame_equal(summary_after.drop(columns=["Result"]), numeric_before)

    next(widget for widget in at.number_input if widget.label == "C-rate").set_value(0.5).run()
    aux_off = auxiliary_override_frame(at)
    assert aux_off.loc[0, "Value"] == pytest.approx(2.25)
    power_before = performance_frame(at)["Net power at POI (MW)"].iloc[0]
    next(widget for widget in at.checkbox if widget.label == "Read flagged aux loads as kW").set_value(
        True
    ).run()
    aux_on = auxiliary_override_frame(at)
    assert aux_on.loc[0, "Value"] == pytest.approx(0.00225)
    assert performance_frame(at)["Net power at POI (MW)"].iloc[0] != pytest.approx(
        power_before
    )

    next(widget for widget in at.number_input if widget.label == "C-rate").set_value(0.25).run()
    at.file_uploader[0].upload(
        "custom.xlsx",
        custom_workbook,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ).run()
    assert not at.exception
    assert any("Uploaded database: 6 container options" in item.value for item in at.caption)
    assert any("8 MWh @ 0.25C" in item.value for item in at.info)
    assert any("28.5 %" in item.value for item in at.caption)

    at.file_uploader[0].upload("products.csv", b"energy,c_rate\n1,0.5", "text/csv").run()
    assert not at.exception
    assert any("CSV files are not supported" in item.value for item in at.error)
    assert any("8 MWh @ 0.25C" in item.value for item in at.info)

    at.file_uploader[0].upload(
        "invalid.xlsx",
        b"not a workbook",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ).run()
    assert not at.exception
    assert any("Last valid database: 6 container options" in item.value for item in at.caption)
    assert any("8 MWh @ 0.25C" in item.value for item in at.info)
    assert any("existing product catalog is unchanged" in item.value for item in at.warning)

    at.file_uploader[0].clear().run()
    assert not at.exception
    assert any("Bundled database: 6 container options" in item.value for item in at.caption)
    assert any("7.044 MWh @ 0.25C" in item.value for item in at.info)
