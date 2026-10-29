"""Catalog validation and isolated persistence tests."""
from __future__ import annotations

from dataclasses import replace
from datetime import datetime, timezone
from decimal import Decimal
import os
import re
from hashlib import sha256
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select, event
from sqlalchemy.schema import CreateSchema, DropSchema
from sqlalchemy.orm import Session
from dotenv import dotenv_values
from sqlalchemy.engine import make_url

from bess import catalog_import
from bess.catalog_import import (
    CatalogIntegrityError,
    CatalogRevisionNotFound,
    CatalogValidationError,
    activate_catalog_revision,
    dry_run_catalog,
    import_catalog,
    load_catalog_revision,
)
from bess.database import load_database
from bess.models import CatalogRevision, ContainerModel, Base


PROJECT_ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = PROJECT_ROOT / "data" / "BESS_DATABASE.xlsx"


@pytest.fixture
def engine():
    database_engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        execution_options={"schema_translate_map": {"bess": None}},
    )

    @event.listens_for(database_engine, "connect")
    def register_now(connection, _record):
        connection.create_function(
            "now", 0, lambda: datetime.now(timezone.utc).isoformat()
        )

    Base.metadata.create_all(database_engine)
    yield database_engine
    database_engine.dispose()


def test_dry_run_reports_hash_counts_warnings_and_does_not_need_database():
    report = dry_run_catalog(WORKBOOK)

    assert report.is_valid
    assert report.source_filename == WORKBOOK.name
    assert report.source_sha256 == sha256(WORKBOOK.read_bytes()).hexdigest()
    assert len(report.containers) == 6
    assert all(item.normalized_ac_parameter_count == 9 for item in report.containers)
    assert all(item.temperature_band_count in (5, 6) for item in report.containers)
    assert any("auxiliary load is missing" in warning for warning in report.warnings)
    assert any("not in database" in warning for warning in report.warnings)


def test_invalid_xlsx_returns_validation_error_without_database_access(tmp_path):
    invalid_file = tmp_path / "invalid.xlsx"
    invalid_file.write_bytes(b"not an xlsx workbook")

    report = dry_run_catalog(invalid_file)

    assert not report.is_valid
    assert report.source_sha256 == sha256(invalid_file.read_bytes()).hexdigest()
    assert any("could not be parsed" in error for error in report.errors)


def test_missing_cycles_source_value_is_not_accepted_as_parser_default(tmp_path):
    changed_workbook = tmp_path / "missing-cycles.xlsx"
    workbook = load_workbook(WORKBOOK)
    worksheet = workbook[workbook.sheetnames[0]]
    for row in worksheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "cycles" in cell.value.lower():
                for offset in (1, 2, 3):
                    worksheet.cell(cell.row, cell.column + offset).value = None
                break
        else:
            continue
        break
    workbook.save(changed_workbook)

    report = dry_run_catalog(changed_workbook)

    assert not report.is_valid
    assert any("source workbook must contain a finite positive cycles value" in error for error in report.errors)


def test_duplicate_source_keys_are_rejected_before_persistence(monkeypatch):
    parsed = load_database(WORKBOOK)
    monkeypatch.setattr(catalog_import, "load_database", lambda _source: [parsed[0], parsed[0]])

    report = dry_run_catalog(WORKBOOK)

    assert not report.is_valid
    assert any("Duplicate container source keys" in error for error in report.errors)


def test_invalid_numeric_parameters_are_rejected(monkeypatch):
    parsed = load_database(WORKBOOK)
    broken = replace(parsed[0], energy_mwh=float("nan"))
    monkeypatch.setattr(catalog_import, "load_database", lambda _source: [broken])

    report = dry_run_catalog(WORKBOOK)

    assert not report.is_valid
    assert any("energy_mwh must be a finite positive number" in error for error in report.errors)


def test_invalid_calendar_rows_are_rejected(monkeypatch):
    parsed = load_database(WORKBOOK)
    broken = replace(parsed[0], cal_rows=[[float("nan")]])
    monkeypatch.setattr(catalog_import, "load_database", lambda _source: [broken])

    report = dry_run_catalog(WORKBOOK)

    assert not report.is_valid
    assert any("does not match temperature bands" in error for error in report.errors)
    assert any("contains an invalid value" in error for error in report.errors)


def test_import_round_trips_every_parser_field_without_activating(engine):
    parsed = {container.id: container for container in load_database(WORKBOOK)}

    with Session(engine) as session:
        result = import_catalog(WORKBOOK, session)
        assert result.created
        assert result.container_count == 6
        session.commit()

    with Session(engine) as session:
        revision = session.get(CatalogRevision, result.revision_id)
        assert revision is not None
        assert not revision.is_active
        restored = {
            container.id: container
            for container in load_catalog_revision(result.revision_id, session)
        }
        assert restored == parsed


def test_repeated_import_is_idempotent(engine):
    with Session(engine) as session:
        first = import_catalog(WORKBOOK, session)
        session.commit()
    with Session(engine) as session:
        second = import_catalog(WORKBOOK, session)
        assert not second.created
        assert second.revision_id == first.revision_id
        assert second.source_sha256 == first.source_sha256
        assert session.scalar(select(func.count()).select_from(CatalogRevision)) == 1
        assert session.scalar(select(func.count()).select_from(ContainerModel)) == 6


def test_import_rolls_back_revision_and_all_container_rows(monkeypatch, engine):
    original_persist = catalog_import._persist_container

    def fail_after_persist(session, revision_id, container):
        original_persist(session, revision_id, container)
        raise RuntimeError("simulated persistence failure")

    monkeypatch.setattr(catalog_import, "_persist_container", fail_after_persist)
    with Session(engine) as session:
        with pytest.raises(RuntimeError, match="simulated persistence failure"):
            import_catalog(WORKBOOK, session)

    with Session(engine) as session:
        assert session.scalar(select(func.count()).select_from(CatalogRevision)) == 0
        assert session.scalar(select(func.count()).select_from(ContainerModel)) == 0


def test_existing_revision_mismatch_is_not_overwritten(engine):
    with Session(engine) as session:
        result = import_catalog(WORKBOOK, session)
        session.commit()

    with Session(engine) as session:
        model = session.scalar(select(ContainerModel))
        assert model is not None
        model.c_rate = model.c_rate + Decimal("0.01")
        changed_source_key = model.source_key
        changed_c_rate = model.c_rate
        session.commit()

    with Session(engine) as session:
        with pytest.raises(CatalogIntegrityError, match="does not match"):
            import_catalog(WORKBOOK, session)

    with Session(engine) as session:
        revision = session.get(CatalogRevision, result.revision_id)
        assert revision is not None
        changed_model = next(
            item for item in revision.container_models if item.source_key == changed_source_key
        )
        assert changed_model.c_rate == changed_c_rate


def test_activation_is_explicit_and_switches_only_one_revision(tmp_path, engine):
    changed_metadata_copy = tmp_path / "same-catalog-new-file-metadata.xlsx"
    workbook = load_workbook(WORKBOOK)
    workbook.properties.creator = "catalog test"
    workbook.save(changed_metadata_copy)

    with Session(engine) as session:
        first = import_catalog(WORKBOOK, session)
        session.commit()
    with Session(engine) as session:
        second = import_catalog(changed_metadata_copy, session)
        session.commit()
    assert first.source_sha256 != second.source_sha256

    with Session(engine) as session:
        assert session.scalar(
            select(func.count()).select_from(CatalogRevision).where(CatalogRevision.is_active.is_(True))
        ) == 0
        session.commit()
        activate_catalog_revision(first.revision_id, session)

    with Session(engine) as session:
        activate_catalog_revision(second.revision_id, session)

    with Session(engine) as session:
        active_ids = session.scalars(
            select(CatalogRevision.id).where(CatalogRevision.is_active.is_(True))
        ).all()
        assert active_ids == [second.revision_id]


def test_unknown_revision_cannot_be_activated(engine):
    from uuid import uuid4

    with Session(engine) as session:
        with pytest.raises(CatalogRevisionNotFound):
            activate_catalog_revision(uuid4(), session)


def test_postgresql_import_reconstructs_the_excel_catalog():
    database_url = os.environ.get("BESS_CATALOG_TEST_DATABASE_URL")
    if not database_url:
        if os.environ.get("BESS_REQUIRE_CATALOG_TESTS") == "1":
            pytest.fail(
                "BESS_REQUIRE_CATALOG_TESTS=1 but BESS_CATALOG_TEST_DATABASE_URL is not configured."
            )
        pytest.skip("Set BESS_CATALOG_TEST_DATABASE_URL to an isolated PostgreSQL test database.")
    try:
        parsed_url = make_url(database_url)
    except Exception:
        pytest.fail("BESS_CATALOG_TEST_DATABASE_URL is not a valid SQLAlchemy URL.")
    target_name = (parsed_url.database or "").casefold()
    if not parsed_url.drivername.startswith("postgresql"):
        pytest.fail("BESS_CATALOG_TEST_DATABASE_URL must use PostgreSQL.")
    if target_name == "bess_db" or not re.search(r"(^|[_-])tests?($|[_-])", target_name):
        pytest.fail(
            "Refusing catalog integration writes unless the explicit database name is test-only and not bess_db."
        )
    application_url = os.environ.get("DATABASE_URL") or dotenv_values(
        PROJECT_ROOT / ".env"
    ).get("DATABASE_URL")
    if application_url:
        try:
            application_database = make_url(application_url).database
        except Exception:
            pytest.fail("Could not safely compare the catalog and application database targets.")
        if application_database == parsed_url.database:
            pytest.fail("Catalog and application database targets must be different.")

    test_schema = f"bess_catalog_test_{uuid4().hex}"
    base_engine = create_engine(parsed_url)
    with base_engine.begin() as connection:
        connection.execute(CreateSchema(test_schema))
    test_engine = base_engine.execution_options(
        schema_translate_map={"bess": test_schema}
    )
    try:
        Base.metadata.create_all(test_engine)
        with Session(test_engine) as session:
            imported = import_catalog(WORKBOOK, session)
            session.commit()
        with Session(test_engine) as session:
            actual = {
                container.id: container
                for container in load_catalog_revision(imported.revision_id, session)
            }
            expected = {container.id: container for container in load_database(WORKBOOK)}
            assert actual == expected
    finally:
        with base_engine.begin() as connection:
            connection.execute(DropSchema(test_schema, cascade=True))
        base_engine.dispose()




